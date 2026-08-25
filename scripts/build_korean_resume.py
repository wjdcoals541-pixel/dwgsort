#!/usr/bin/env python3
"""Build a polished two-page Korean resume from the supplied master profile.

The source PDF contains narrative notes and AI-facing instructions. This builder
uses only resume-relevant biographical and career facts; it intentionally omits
instructional text, legal/payment details, and unrelated private context.
"""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "output" / "documents" / "정채민_국문이력서_2026.docx"
PHOTO_PATH = Path(r"C:\Users\jcm07\Downloads\DUB00327 copya copy-정채민.jpg")

FONT_KR = "맑은 고딕"
FONT_LATIN = "Arial"

NAVY = "17324D"
TEAL = "2E6F73"
INK = "1E252B"
MUTED = "66727C"
LIGHT_BLUE = "EEF4F7"
LIGHT_GRAY = "F6F8F9"
HAIRLINE = "D7E0E6"
WHITE = "FFFFFF"

CELL_MARGINS = {"top": 75, "bottom": 75, "start": 120, "end": 120}
TABLE_INDENT_DXA = CELL_MARGINS["start"]


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    color: str | None = None,
    italic: bool | None = None,
    latin_font: str = FONT_LATIN,
    east_asia_font: str = FONT_KR,
) -> None:
    run.font.name = latin_font
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin_font)
    r_fonts.set(qn("w:hAnsi"), latin_font)
    r_fonts.set(qn("w:cs"), latin_font)
    r_fonts.set(qn("w:eastAsia"), east_asia_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = rgb(color)
    if italic is not None:
        run.italic = italic


def set_style_font(style, *, size: float, bold: bool = False, color: str = INK) -> None:
    style.font.name = FONT_LATIN
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = rgb(color)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_fonts.set(qn("w:cs"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), FONT_KR)


def ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_paragraph_border_bottom(paragraph, *, color: str = HAIRLINE, size: int = 8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = ensure_child(p_pr, "w:pBdr")
    bottom = ensure_child(p_bdr, "w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = ensure_child(p_pr, "w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = ensure_child(tc_pr, "w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, **borders) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = ensure_child(tc_pr, "w:tcBorders")
    for edge, values in borders.items():
        element = ensure_child(tc_borders, f"w:{edge}")
        for key, value in values.items():
            element.set(qn(f"w:{key}"), str(value))


def clear_cell_borders(cell) -> None:
    none = {"val": "nil"}
    set_cell_borders(
        cell,
        top=none,
        left=none,
        bottom=none,
        right=none,
        insideH=none,
        insideV=none,
    )


def set_cell_margins(cell, margins: dict[str, int] = CELL_MARGINS) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = ensure_child(tc_pr, "w:tcMar")
    for side in ("top", "bottom", "start", "end"):
        margin = ensure_child(tc_mar, f"w:{side}")
        margin.set(qn("w:w"), str(margins[side]))
        margin.set(qn("w:type"), "dxa")


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_row_min_height(row, height_dxa: int) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_height = ensure_child(tr_pr, "w:trHeight")
    tr_height.set(qn("w:val"), str(height_dxa))
    tr_height.set(qn("w:hRule"), "atLeast")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def apply_table_geometry(
    table,
    widths_dxa: list[int],
    *,
    indent_dxa: int = TABLE_INDENT_DXA,
    margins: dict[str, int] = CELL_MARGINS,
) -> None:
    if not widths_dxa or any(width <= 0 for width in widths_dxa):
        raise ValueError(f"Invalid table widths: {widths_dxa}")
    if any(len(row.cells) != len(widths_dxa) for row in table.rows):
        raise ValueError("Merged or irregular rows are not supported by this geometry helper")

    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = ensure_child(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))

    tbl_ind = ensure_child(tbl_pr, "w:tblInd")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))

    layout = ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for col_idx, width in enumerate(widths_dxa):
        table.columns[col_idx].width = Twips(width)

    for row in table.rows:
        row.height = None
        set_row_cant_split(row)
        for col_idx, cell in enumerate(row.cells):
            width = widths_dxa[col_idx]
            cell.width = Twips(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = ensure_child(tc_pr, "w:tcW")
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            set_cell_margins(cell, margins)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph, field_code: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {field_code} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run = paragraph.add_run()
    set_run_font(run, size=8, color=MUTED)
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(result)
    run._r.append(end)


def add_bullet_numbering(doc: Document, *, left_dxa: int = 480, hanging_dxa: int = 240) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(el.get(qn("w:abstractNumId")))
        for el in numbering.findall(qn("w:abstractNum"))
        if el.get(qn("w:abstractNumId")) is not None
    ]
    num_ids = [
        int(el.get(qn("w:numId")))
        for el in numbering.findall(qn("w:num"))
        if el.get(qn("w:numId")) is not None
    ]
    abstract_id = (max(abstract_ids) + 1) if abstract_ids else 0
    num_id = (max(num_ids) + 1) if num_ids else 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.extend([start, num_fmt, lvl_text, lvl_jc])

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left_dxa))
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left_dxa))
    ind.set(qn("w:hanging"), str(hanging_dxa))
    p_pr.extend([tabs, ind])
    lvl.append(p_pr)

    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), FONT_KR)
    r_pr.append(r_fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = ensure_child(num_pr, "w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = ensure_child(num_pr, "w:numId")
    num_id_el.set(qn("w:val"), str(num_id))


def add_bullet(doc: Document, text: str, num_id: int):
    paragraph = doc.add_paragraph(style="Resume Bullet")
    apply_numbering(paragraph, num_id)
    run = paragraph.add_run(text)
    set_run_font(run, size=9.15, color=INK)
    return paragraph


def add_cell_paragraph(cell, text: str = "", *, style: str | None = None):
    paragraph = cell.paragraphs[0] if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text else cell.add_paragraph()
    if style:
        paragraph.style = style
    if text:
        run = paragraph.add_run(text)
        set_run_font(run, size=9.0, color=INK)
    return paragraph


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, size=9.3, color=INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    heading_1 = styles["Heading 1"]
    set_style_font(heading_1, size=11.6, bold=True, color=NAVY)
    heading_1.paragraph_format.space_before = Pt(8)
    heading_1.paragraph_format.space_after = Pt(4.5)
    heading_1.paragraph_format.line_spacing = 1.0
    heading_1.paragraph_format.keep_with_next = True
    heading_1.paragraph_format.keep_together = True

    heading_2 = styles["Heading 2"]
    set_style_font(heading_2, size=10.2, bold=True, color=TEAL)
    heading_2.paragraph_format.space_before = Pt(5)
    heading_2.paragraph_format.space_after = Pt(2)
    heading_2.paragraph_format.line_spacing = 1.0
    heading_2.paragraph_format.keep_with_next = True
    heading_2.paragraph_format.keep_together = True

    custom = {
        "Resume Bullet": (9.15, False, INK, 0, 1.5, 1.08),
        "Resume Meta": (8.7, False, MUTED, 0, 1.5, 1.0),
        "Resume Company": (10.25, True, NAVY, 5, 2.5, 1.0),
        "Resume Small": (8.4, False, MUTED, 0, 2, 1.05),
        "Resume Table": (8.8, False, INK, 0, 0, 1.05),
        "Resume Label": (8.65, True, NAVY, 0, 0, 1.0),
    }
    for name, (size, bold, color, before, after, line) in custom.items():
        style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(style, size=size, bold=bold, color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = line
        style.paragraph_format.widow_control = True
        if name in {"Resume Company", "Resume Label"}:
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.keep_together = True


def add_section_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text, style="Heading 1")
    set_paragraph_border_bottom(paragraph, color=TEAL, size=10)


def add_company_heading(doc: Document, company: str, role: str, period: str, right_edge_dxa: int) -> None:
    paragraph = doc.add_paragraph(style="Resume Company")
    paragraph.paragraph_format.tab_stops.add_tab_stop(Twips(right_edge_dxa), WD_TAB_ALIGNMENT.RIGHT)
    company_run = paragraph.add_run(company)
    set_run_font(company_run, size=10.4, bold=True, color=NAVY)
    role_run = paragraph.add_run(f"  |  {role}")
    set_run_font(role_run, size=9.0, bold=False, color=MUTED)
    date_run = paragraph.add_run(f"\t{period}")
    set_run_font(date_run, size=8.8, color=MUTED)


def add_labeled_line(doc: Document, label: str, value: str) -> None:
    paragraph = doc.add_paragraph(style="Resume Small")
    paragraph.paragraph_format.left_indent = Twips(260)
    label_run = paragraph.add_run(f"{label} | ")
    set_run_font(label_run, size=8.4, bold=True, color=TEAL)
    value_run = paragraph.add_run(value)
    set_run_font(value_run, size=8.4, color=MUTED)


def set_table_cell_text(
    cell,
    text: str,
    *,
    size: float = 8.8,
    bold: bool = False,
    color: str = INK,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    style: str = "Resume Table",
) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.style = style
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_key_value_table(doc: Document, rows: list[tuple[str, str]], table_width_dxa: int) -> None:
    label_width = 1520
    value_width = table_width_dxa - label_width
    table = doc.add_table(rows=len(rows), cols=2)
    apply_table_geometry(table, [label_width, value_width])
    for index, (label, value) in enumerate(rows):
        row = table.rows[index]
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        set_cell_shading(row.cells[1], WHITE)
        set_table_cell_text(row.cells[0], label, size=8.65, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_table_cell_text(row.cells[1], value, size=8.75, color=INK)
        for cell in row.cells:
            set_cell_borders(
                cell,
                top={"val": "single", "sz": "4", "color": HAIRLINE},
                left={"val": "single", "sz": "4", "color": HAIRLINE},
                bottom={"val": "single", "sz": "4", "color": HAIRLINE},
                right={"val": "single", "sz": "4", "color": HAIRLINE},
            )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 0.4


def build_resume(output_path: Path) -> None:
    if not PHOTO_PATH.exists():
        raise FileNotFoundError(f"Resume photo not found: {PHOTO_PATH}")

    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(13.5)
    section.bottom_margin = Mm(12.5)
    section.left_margin = Mm(16)
    section.right_margin = Mm(16)
    section.header_distance = Mm(6)
    section.footer_distance = Mm(7)

    content_width_dxa = int(
        round(
            section.page_width.twips
            - section.left_margin.twips
            - section.right_margin.twips
        )
    )
    table_width_dxa = content_width_dxa - TABLE_INDENT_DXA

    doc.core_properties.title = "정채민 국문 이력서"
    doc.core_properties.subject = "생산기술·설비·CS·기계기술 지원용 국문 이력서"
    doc.core_properties.author = "정채민"
    doc.core_properties.keywords = "국문 이력서, 현장형 엔지니어, 생산기술, 설비, CS"
    doc.core_properties.comments = "2026-08-16 마스터 프로필 기준 작성"

    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    bullet_num_id = add_bullet_numbering(doc)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer_paragraph.paragraph_format.tab_stops.add_tab_stop(
        Twips(content_width_dxa), WD_TAB_ALIGNMENT.RIGHT
    )
    footer_paragraph.paragraph_format.space_before = Pt(2)
    footer_paragraph.paragraph_format.space_after = Pt(0)
    set_paragraph_border_bottom(footer_paragraph, color=HAIRLINE, size=4)
    left_footer = footer_paragraph.add_run("정채민 | 국문 이력서")
    set_run_font(left_footer, size=8, color=MUTED)
    tab_run = footer_paragraph.add_run("\t")
    set_run_font(tab_run, size=8, color=MUTED)
    add_page_field(footer_paragraph, "PAGE")
    slash = footer_paragraph.add_run(" / ")
    set_run_font(slash, size=8, color=MUTED)
    add_page_field(footer_paragraph, "NUMPAGES")

    # First-page identity block with the user-supplied identification photo.
    header_table = doc.add_table(rows=1, cols=2)
    photo_width = 1840
    apply_table_geometry(header_table, [table_width_dxa - photo_width, photo_width])
    set_row_min_height(header_table.rows[0], 2200)
    left_cell, photo_cell = header_table.rows[0].cells
    clear_cell_borders(left_cell)
    clear_cell_borders(photo_cell)
    set_cell_borders(
        photo_cell,
        top={"val": "single", "sz": "7", "color": HAIRLINE},
        left={"val": "single", "sz": "7", "color": HAIRLINE},
        bottom={"val": "single", "sz": "7", "color": HAIRLINE},
        right={"val": "single", "sz": "7", "color": HAIRLINE},
    )
    set_cell_shading(photo_cell, LIGHT_GRAY)
    photo_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    name_p = left_cell.paragraphs[0]
    name_p.paragraph_format.space_before = Pt(0)
    name_p.paragraph_format.space_after = Pt(2)
    name_run = name_p.add_run("정채민")
    set_run_font(name_run, size=25, bold=True, color=NAVY)

    role_p = left_cell.add_paragraph()
    role_p.paragraph_format.space_after = Pt(4)
    role_run = role_p.add_run("현장형 엔지니어")
    set_run_font(role_run, size=11.2, bold=True, color=TEAL)
    field_run = role_p.add_run("  |  생산기술 · 설비/유지보수 · CS · 기계기술")
    set_run_font(field_run, size=9.0, color=MUTED)

    info_rows = [
        (("연락처", "010-9668-4878"), ("이메일", "wjdcoals541@gmail.com")),
        (("거주지", "인천광역시 부평구"), ("생년월일", "1998.07.31 / 남")),
    ]
    for left_info, right_info in info_rows:
        p = left_cell.add_paragraph(style="Resume Meta")
        p.paragraph_format.tab_stops.add_tab_stop(Twips(4100), WD_TAB_ALIGNMENT.LEFT)
        label_run = p.add_run(f"{left_info[0]}  ")
        set_run_font(label_run, size=8.7, bold=True, color=NAVY)
        value_run = p.add_run(left_info[1])
        set_run_font(value_run, size=8.7, color=INK)
        p.add_run("\t")
        label_run_2 = p.add_run(f"{right_info[0]}  ")
        set_run_font(label_run_2, size=8.7, bold=True, color=NAVY)
        value_run_2 = p.add_run(right_info[1])
        set_run_font(value_run_2, size=8.7, color=INK)

    photo_p = photo_cell.paragraphs[0]
    photo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    photo_p.paragraph_format.space_after = Pt(0)
    photo_run = photo_p.add_run()
    picture = photo_run.add_picture(str(PHOTO_PATH), width=Mm(27.5))
    picture._inline.docPr.set("title", "증명사진")
    picture._inline.docPr.set("descr", "정채민 증명사진")

    add_section_heading(doc, "직무 요약")
    summary_table = doc.add_table(rows=1, cols=1)
    apply_table_geometry(summary_table, [table_width_dxa])
    set_cell_shading(summary_table.cell(0, 0), LIGHT_BLUE)
    set_cell_borders(
        summary_table.cell(0, 0),
        top={"val": "single", "sz": "6", "color": HAIRLINE},
        left={"val": "single", "sz": "10", "color": TEAL},
        bottom={"val": "single", "sz": "6", "color": HAIRLINE},
        right={"val": "single", "sz": "6", "color": HAIRLINE},
    )
    summary = (
        "자동화 기계 조립·가공·배선과 관로 CAD·수충격 해석을 경험했습니다. "
        "도면과 실제 부품·자료의 불일치를 확인해 현장 재가공, 연구팀 및 외부업체 협의로 해결했으며, "
        "반복적인 CAD 데이터 정리 업무는 Python 기반 도구로 개선해 프로젝트에 따라 하루 이상 걸리던 작업을 약 10분 수준으로 줄였습니다."
    )
    set_table_cell_text(summary_table.cell(0, 0), summary, size=9.15, color=INK)

    add_section_heading(doc, "핵심 역량")
    skill_items = [
        "자동화기계 조립·시운전",
        "도면 해독·현장 수정",
        "Pipe2012 수충격 해석",
        "CAD·기술자료 검토",
        "기술보고서·대외협의",
        "Python 업무자동화",
    ]
    skills_table = doc.add_table(rows=2, cols=3)
    skill_widths = [table_width_dxa // 3, table_width_dxa // 3]
    skill_widths.append(table_width_dxa - sum(skill_widths))
    apply_table_geometry(skills_table, skill_widths)
    for index, text in enumerate(skill_items):
        cell = skills_table.rows[index // 3].cells[index % 3]
        set_cell_shading(cell, LIGHT_GRAY)
        set_cell_borders(
            cell,
            top={"val": "single", "sz": "4", "color": HAIRLINE},
            left={"val": "single", "sz": "4", "color": HAIRLINE},
            bottom={"val": "single", "sz": "4", "color": HAIRLINE},
            right={"val": "single", "sz": "4", "color": HAIRLINE},
        )
        set_table_cell_text(cell, text, size=8.8, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)

    add_section_heading(doc, "경력 요약")
    career_rows = [
        ("2025.12.18 - 2026.06.21", "플로우테크(주) / 주임", "Pipe2012 수충격 해석, CAD·수리계산서 검토, 기술보고서 및 대외 기술협의"),
        ("2022.02.14 - 2024.04.10", "(주)리팩 / 사원", "로터리 자동 포장기계 조립, 부품 가공·용접, 유압 및 기본 전기배선"),
        ("2017.10 - 2017.12", "(주)엘리비젼 / 현장실습", "키오스크 제작·부품 조립, 품질 확인, 현장 점검 및 A/S 보조"),
    ]
    career_table = doc.add_table(rows=1 + len(career_rows), cols=3)
    career_widths = [2100, 2400, table_width_dxa - 4500]
    apply_table_geometry(career_table, career_widths)
    headers = ["기간", "회사 / 직급", "주요 업무"]
    for idx, header in enumerate(headers):
        cell = career_table.rows[0].cells[idx]
        set_cell_shading(cell, NAVY)
        set_table_cell_text(cell, header, size=8.6, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_borders(
            cell,
            top={"val": "single", "sz": "5", "color": NAVY},
            left={"val": "single", "sz": "5", "color": NAVY},
            bottom={"val": "single", "sz": "5", "color": NAVY},
            right={"val": "single", "sz": "5", "color": NAVY},
        )
    set_repeat_table_header(career_table.rows[0])
    for row_index, values in enumerate(career_rows, start=1):
        for col_index, value in enumerate(values):
            cell = career_table.rows[row_index].cells[col_index]
            set_cell_shading(cell, WHITE if row_index % 2 else LIGHT_GRAY)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_index < 2 else WD_ALIGN_PARAGRAPH.LEFT
            set_table_cell_text(cell, value, size=8.45, color=INK, align=align)
            set_cell_borders(
                cell,
                top={"val": "single", "sz": "4", "color": HAIRLINE},
                left={"val": "single", "sz": "4", "color": HAIRLINE},
                bottom={"val": "single", "sz": "4", "color": HAIRLINE},
                right={"val": "single", "sz": "4", "color": HAIRLINE},
            )

    add_section_heading(doc, "주요 경력")
    add_company_heading(doc, "플로우테크(주)", "기술영업팀 주임", "2025.12.18 - 2026.06.21", content_width_dxa)
    for text in [
        "Pipe2012(P2K)로 펌프 H-Q 곡선·RPM·밸브 개폐 시퀀스를 반영한 수충격 해석 수행",
        "관로 CAD 종단면도, 수리계산서, 기계실 도면, 펌프 기술자료의 수치·조건 정합성 검토",
        "엔지니어링사·시공사와 누락 자료 및 기준을 확인하고 발주처 기준에 맞춰 기술보고서 작성",
        "재직 중 31건 이상의 프로젝트에서 해석·자료검토·보고서 업무 수행",
    ]:
        add_bullet(doc, text, bullet_num_id)
    add_labeled_line(
        doc,
        "대표 프로젝트",
        "청주 도시물길 조성사업 실시설계 · 서울시 암사/뚝도 정수센터 송수관 이중화 공사 · 과천 공공하수처리시설 현대화 수충격 해석(TK)",
    )

    add_company_heading(doc, "(주)리팩", "조립1팀 사원", "2022.02.14 - 2024.04.10", content_width_dxa)
    for text in [
        "조립도와 부품도면을 바탕으로 로터리 자동 포장기계 조립 및 시운전",
        "드릴링머신·절단기·그라인더 가공, 용접, 유압배선 및 기본 전기배선 수행",
        "도면과 실물 치수 불일치 또는 간섭 발생 시 위치·치수 확인 후 현장 재가공",
        "설계 변경이 필요한 사안은 연구팀에 실제 치수와 필요한 수정치를 전달해 해결방안 협의",
    ]:
        add_bullet(doc, text, bullet_num_id)

    # Deliberate two-page structure.
    page_break = doc.add_paragraph()
    page_break.paragraph_format.space_after = Pt(0)
    page_break.add_run().add_break(WD_BREAK.PAGE)

    add_section_heading(doc, "경력사항 (계속)")
    add_company_heading(doc, "(주)엘리비젼", "현장실습생(인턴)", "2017.10 - 2017.12", content_width_dxa)
    for text in [
        "학과장 추천으로 참여한 현장실습에서 키오스크 제작 및 부품 조립 수행",
        "완제품의 외관·조립상태·먼지·흠집을 확인하며 품질 기준과 작업순서의 중요성 학습",
        "설치 현장 점검과 A/S를 보조하고 이상 부위 확인 및 부품 교체 경험",
    ]:
        add_bullet(doc, text, bullet_num_id)

    add_section_heading(doc, "주요 업무개선 프로젝트")
    project_title = doc.add_paragraph(style="Heading 2")
    title_run = project_title.add_run("DWGSort | CAD 데이터 변환 업무개선")
    set_run_font(title_run, size=10.2, bold=True, color=TEAL)
    add_key_value_table(
        doc,
        [
            ("과제", "CAD 종단면도의 누가거리·관저고 데이터를 Excel 및 수충격 해석 입력자료로 옮기는 반복 작업 개선"),
            ("실행", "AutoCAD 데이터 추출 기능과 Python을 결합하고, 실제 도면에서 발생한 오류를 반복 확인하며 변환 로직 보완"),
            ("결과", "프로젝트에 따라 하루 이상 걸리던 작업을 약 10분 수준으로 단축하고 실제 업무에 지속 활용"),
            ("도구", "Python, PySide6, pandas, openpyxl, PyMuPDF, AutoCAD 데이터 추출"),
        ],
        table_width_dxa,
    )

    add_section_heading(doc, "학력")
    education_table = doc.add_table(rows=1, cols=2)
    education_widths = [2050, table_width_dxa - 2050]
    apply_table_geometry(education_table, education_widths)
    set_cell_shading(education_table.cell(0, 0), LIGHT_BLUE)
    set_cell_shading(education_table.cell(0, 1), WHITE)
    set_table_cell_text(
        education_table.cell(0, 0),
        "2017.02.27 - 2022.02.11",
        size=8.6,
        bold=True,
        color=NAVY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    education_cell = education_table.cell(0, 1)
    education_cell.paragraphs[0].clear()
    p = education_cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    school = p.add_run("인하공업전문대학 메카트로닉스과  |  공학전문학사")
    set_run_font(school, size=9.0, bold=True, color=INK)
    p2 = education_cell.add_paragraph(style="Resume Small")
    detail = p2.add_run("총 80학점 · 평점 3.44/4.5 · 주요 학습: C언어, PLC, 유압/공기압제어, 도면해독, 3D CAD, 로봇공학, 마이크로프로세서")
    set_run_font(detail, size=8.35, color=MUTED)
    for cell in education_table.rows[0].cells:
        set_cell_borders(
            cell,
            top={"val": "single", "sz": "4", "color": HAIRLINE},
            left={"val": "single", "sz": "4", "color": HAIRLINE},
            bottom={"val": "single", "sz": "4", "color": HAIRLINE},
            right={"val": "single", "sz": "4", "color": HAIRLINE},
        )

    add_section_heading(doc, "기술 및 역량")
    add_key_value_table(
        doc,
        [
            ("현장 제조", "자동화 장비 조립, 부품 가공, 드릴링·절단·그라인딩, 용접, 유압배선, 기본 전기배선, 시운전·수정"),
            ("도면·해석", "도면 해독, AutoCAD, 관로 종단면도·수리계산서·기계실 도면 검토, Pipe2012 수충격 해석"),
            ("제어 기초", "PLC, 유압·공기압제어, 전기전자, 모터제어, HMI, 계측센서 - 대학 실습 기반"),
            ("업무자동화", "Python, PySide6, pandas, openpyxl, PyMuPDF - 반복 데이터 변환 프로그램 제작·개선"),
            ("문서·협업", "기술보고서 작성, 엔지니어링사·시공사·연구팀과 기술조건·치수·자료·일정 협의"),
        ],
        table_width_dxa,
    )

    add_section_heading(doc, "자격 및 병역")
    credential_table = doc.add_table(rows=2, cols=2)
    credential_widths = [1850, table_width_dxa - 1850]
    apply_table_geometry(credential_table, credential_widths)
    credentials = [
        ("운전면허", "자동차운전면허 제1종 보통 · 최초교부 2017.01.18"),
        ("병역", "육군 포병 병장 전역 · 2018.09.11 - 2020.04.26 · 특급전사 · 포병전술경연대회 K55A1 자주포 사수 선발·참가"),
    ]
    for row_index, (label, value) in enumerate(credentials):
        set_cell_shading(credential_table.rows[row_index].cells[0], LIGHT_BLUE)
        set_table_cell_text(
            credential_table.rows[row_index].cells[0],
            label,
            size=8.65,
            bold=True,
            color=NAVY,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_table_cell_text(credential_table.rows[row_index].cells[1], value, size=8.75, color=INK)
        for cell in credential_table.rows[row_index].cells:
            set_cell_borders(
                cell,
                top={"val": "single", "sz": "4", "color": HAIRLINE},
                left={"val": "single", "sz": "4", "color": HAIRLINE},
                bottom={"val": "single", "sz": "4", "color": HAIRLINE},
                right={"val": "single", "sz": "4", "color": HAIRLINE},
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def validate_resume(path: Path) -> None:
    if not path.exists() or path.stat().st_size < 10_000:
        raise AssertionError(f"DOCX was not created correctly: {path}")

    doc = Document(path)
    text_parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            text_parts.extend(cell.text for cell in row.cells)
    full_text = "\n".join(text_parts)

    required = [
        "정채민",
        "010-9668-4878",
        "wjdcoals541@gmail.com",
        "플로우테크(주)",
        "(주)리팩",
        "(주)엘리비젼",
        "Pipe2012(P2K)",
        "31건 이상의 프로젝트",
        "약 10분 수준으로 단축",
        "인하공업전문대학 메카트로닉스과",
        "자동차운전면허 제1종 보통",
    ]
    missing = [item for item in required if item not in full_text]
    if missing:
        raise AssertionError(f"Required resume content missing: {missing}")

    forbidden = [
        "AI MASTER PROFILE",
        "다른 AI에게",
        "임금체불",
        "체불임금",
        "법적 분쟁",
        "프로그램 소유권",
        "게임을 즐기며",
        "Codex",
        "Claude Code",
    ]
    found_forbidden = [item for item in forbidden if item in full_text]
    if found_forbidden:
        raise AssertionError(f"Private or instructional content leaked into resume: {found_forbidden}")

    section = doc.sections[0]
    page_mm = (section.page_width.mm, section.page_height.mm)
    if not (209.5 <= page_mm[0] <= 210.5 and 296.5 <= page_mm[1] <= 297.5):
        raise AssertionError(f"Expected A4 page geometry, got {page_mm}")

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
        numbering_root = ET.fromstring(archive.read("word/numbering.xml"))
        media_files = [name for name in archive.namelist() if name.startswith("word/media/")]
    attr = lambda name: f"{{{ns['w']}}}{name}"

    tables = root.findall(".//w:tbl", ns)
    if len(tables) != 8:
        raise AssertionError(f"Expected 8 body tables, found {len(tables)}")
    for table_index, tbl in enumerate(tables, start=1):
        tbl_w = tbl.find("w:tblPr/w:tblW", ns)
        tbl_ind = tbl.find("w:tblPr/w:tblInd", ns)
        grid = [int(col.get(attr("w"), "0")) for col in tbl.findall("w:tblGrid/w:gridCol", ns)]
        if tbl_w is None or tbl_w.get(attr("type")) != "dxa":
            raise AssertionError(f"Table {table_index} lacks DXA width")
        table_width = int(tbl_w.get(attr("w"), "0"))
        if sum(grid) != table_width:
            raise AssertionError(f"Table {table_index} grid width mismatch")
        if tbl_ind is None or int(tbl_ind.get(attr("w"), "-1")) != TABLE_INDENT_DXA:
            raise AssertionError(f"Table {table_index} indent mismatch")
        for row_index, row in enumerate(tbl.findall("w:tr", ns), start=1):
            cell_widths = []
            for cell in row.findall("w:tc", ns):
                tc_w = cell.find("w:tcPr/w:tcW", ns)
                cell_widths.append(int(tc_w.get(attr("w"), "0")) if tc_w is not None else 0)
            if cell_widths != grid:
                raise AssertionError(
                    f"Table {table_index} row {row_index} cell widths {cell_widths} do not match {grid}"
                )

    bullet_levels = numbering_root.findall(".//w:numFmt[@w:val='bullet']", ns)
    numbered_paragraphs = root.findall(".//w:pPr/w:numPr", ns)
    if not bullet_levels or len(numbered_paragraphs) < 10:
        raise AssertionError("Real bullet numbering was not encoded")

    page_breaks = root.findall(".//w:br[@w:type='page']", ns)
    if len(page_breaks) != 1:
        raise AssertionError(f"Expected one deliberate page break, found {len(page_breaks)}")
    if len(media_files) != 1:
        raise AssertionError(f"Expected one embedded resume photo, found {media_files}")

    drawing_properties = root.findall(".//w:drawing//wp:docPr", {
        **ns,
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    })
    if not any(item.get("descr") == "정채민 증명사진" for item in drawing_properties):
        raise AssertionError("Embedded resume photo is missing descriptive alt text")

    print(f"Created: {path}")
    print(f"Size: {path.stat().st_size} bytes")
    print(f"Tables: {len(tables)} (DXA geometry verified)")
    print(f"Bullet paragraphs: {len(numbered_paragraphs)}")
    print(f"Embedded photos: {len(media_files)} (alt text verified)")
    print("A4 geometry: verified")
    print("Required content: verified")
    print("Instructional/private exclusions: verified")


def main() -> None:
    build_resume(OUTPUT_PATH)
    validate_resume(OUTPUT_PATH)


if __name__ == "__main__":
    main()
